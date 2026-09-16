"""Group 4b — export endpoints (GET /api/jobs/{id}/export/{json|csv|excel}).

Jobs + posts are seeded directly through the models (helpers.insert_completed_job)
so these tests verify the export WIRING (build_export -> exporters ->
FileResponse, content types, 404/409 envelopes) without depending on the
background worker.
"""
from __future__ import annotations

import io
import json

from helpers import error_envelope, insert_completed_job, sample_posts

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _post_job_and_export(client, fmt: str, posts):
    job_id = insert_completed_job(posts)
    resp = client.get(f"/api/jobs/{job_id}/export/{fmt}")
    assert resp.status_code == 200, resp.text[:500]
    return resp


def test_export_unknown_job_404(authed_client):
    for fmt in ("json", "csv", "excel"):
        resp = authed_client.get(f"/api/jobs/does-not-exist/export/{fmt}")
        assert resp.status_code == 404
        err = error_envelope(resp.json())
        assert err["code"] == "not_found"


def test_export_json_endpoint(authed_client):
    posts = sample_posts(3)
    resp = _post_job_and_export(authed_client, "json", posts)
    assert resp.headers["content-type"].startswith("application/json")
    loaded = json.loads(resp.content)
    assert len(loaded) == 3
    # every stored post round-trips (scraped_at is added by serialization,
    # published_at is normalized from +00:00 to Z by iso_format)
    for original, exported in zip(posts, loaded):
        for key in (
            "post_id", "post_url", "text", "post_type",
            "timestamp", "hashtags",
        ):
            assert exported[key] == original[key], f"{key} drifted on export"
        # published_at normalizes +00:00 -> Z via serialization.iso_format
        assert exported["published_at"] is not None


def test_export_csv_endpoint(authed_client):
    posts = sample_posts(3)
    resp = _post_job_and_export(authed_client, "csv", posts)
    assert resp.headers["content-type"].startswith("text/csv")
    assert resp.content[:3] == b"\xef\xbb\xbf", "CSV must carry the UTF-8 BOM"


def test_export_excel_endpoint(authed_client):
    posts = sample_posts(3)
    resp = _post_job_and_export(authed_client, "excel", posts)
    assert resp.headers["content-type"].startswith(XLSX_MEDIA)
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Posts", "Engagement", "Media", "Metadata"]
    assert wb["Posts"].max_row == 1 + 3  # header + 3 posts


def test_export_running_job_conflict(authed_client):
    job_id = insert_completed_job(sample_posts(1), status="running")
    resp = authed_client.get(f"/api/jobs/{job_id}/export/json")
    assert resp.status_code == 409
    err = error_envelope(resp.json())
    assert err["code"] == "job_running"


def test_export_queued_job_conflict(authed_client):
    job_id = insert_completed_job(sample_posts(1), status="queued")
    resp = authed_client.get(f"/api/jobs/{job_id}/export/csv")
    assert resp.status_code == 409


def test_export_invalid_format_400(authed_client):
    job_id = insert_completed_job(sample_posts(1))
    resp = authed_client.get(f"/api/jobs/{job_id}/export/yaml")
    assert resp.status_code == 400  # FastAPI Literal rejects it
    err = error_envelope(resp.json())
    assert err["code"] == "validation_error"